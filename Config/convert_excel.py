import os
import sys
import openpyxl
from pathlib import Path
import argparse

class ExcelToRobloxConverter:
    def __init__(self, input_dir, output_dir):
        self.input_dir = input_dir
        self.output_dir = output_dir
        
        # 确保输出目录存在
        os.makedirs(output_dir, exist_ok=True)
        
    def convert_type(self, value, type_name):
        """根据类型名转换值"""
        # 处理空值
        if value is None:
            if type_name == "string":
                return '""'
            elif type_name in ["int", "int[]", "int[][]"]:
                return "0"
            elif type_name in ["float", "float[]", "float[][]"]:
                return "0.0"
            else:
                return "nil"
        
        # 处理数组类型
        if type_name.endswith("[]") or type_name.endswith("[][]"):
            base_type = type_name.replace("[]", "").replace("[]", "")
            
            # 处理一维数组
            if type_name.endswith("[]") and not type_name.endswith("[][]"):
                if isinstance(value, str):
                    # 尝试将字符串分割成数组
                    items = [item.strip() for item in value.split(',')] if ',' in value else [value]
                    converted_items = [self.convert_type(item, base_type) for item in items]
                    return "{" + ", ".join(converted_items) + "}"
                elif isinstance(value, (list, tuple)):
                    converted_items = [self.convert_type(item, base_type) for item in value]
                    return "{" + ", ".join(converted_items) + "}"
                else:
                    # 单个值当作数组处理
                    return "{" + self.convert_type(value, base_type) + "}"
            
            # 处理二维数组
            elif type_name.endswith("[][]"):
                if isinstance(value, str):
                    # 尝试将字符串分割成二维数组
                    rows = [row.strip() for row in value.split(';')] if ';' in value else [value]
                    result = []
                    for row in rows:
                        items = [item.strip() for item in row.split(',')] if ',' in row else [row]
                        converted_items = [self.convert_type(item, base_type) for item in items]
                        result.append("{" + ", ".join(converted_items) + "}")
                    return "{" + ", ".join(result) + "}"
                else:
                    # 默认处理
                    return "{" + self.convert_type(value, base_type) + "}"
        
        # 处理基本类型
        if type_name == "string":
            # 确保字符串中的特殊字符被正确转义
            if isinstance(value, str):
                value = value.replace('"', '\\"').replace("\n", "\\n")
            return f'"{value}"'
        elif type_name == "int":
            try:
                return str(int(float(value)))
            except (ValueError, TypeError):
                return "0"
        elif type_name == "float":
            try:
                return str(float(value))
            except (ValueError, TypeError):
                return "0.0"
        elif type_name == "bool":
            if isinstance(value, bool):
                return str(value).lower()
            elif isinstance(value, str):
                return str(value.lower() in ["true", "1", "yes", "y"]).lower()
            else:
                return str(bool(value)).lower()
        else:
            # 默认作为字符串处理
            return f'"{value}"'
            
    def process_excel_file(self, file_path):
        """处理单个Excel文件"""
        print(f"处理文件: {file_path}...")
        
        # 读取Excel文件
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheet = wb.active
        
        # 获取表格名称(文件名，不含扩展名)
        table_name = Path(file_path).stem
        
        # 读取数据类型行（第3行，索引为2）
        types = []
        for col in range(1, sheet.max_column + 1):
            cell_value = sheet.cell(row=3, column=col).value
            types.append(str(cell_value).lower() if cell_value else "string")
        
        # 读取字段名（第1行，索引为0）
        field_names = []
        for col in range(1, sheet.max_column + 1):
            cell_value = sheet.cell(row=1, column=col).value
            if col == 1:  # 第一列固定为onlyID
                field_names.append("onlyID")
            else:
                field_names.append(str(cell_value) if cell_value else f"Column{col}")
        
        # 创建Lua代码
        lua_code = f"-- 自动生成的数据模块: {table_name}\n"
        lua_code += f"local {table_name} = {{\n"
        
        # 存储所有数据行
        all_rows = []
        
        # 从第6行开始读取数据（索引为5，跳过表头、类型、分组行和中文提示行）
        for row_idx in range(6, sheet.max_row + 1):
            row_data = {}
            row_values = []
            
            # 检查行是否为空
            is_empty_row = True
            for col in range(1, sheet.max_column + 1):
                cell_value = sheet.cell(row=row_idx, column=col).value
                if cell_value is not None:
                    is_empty_row = False
                    break
                    
            if is_empty_row:
                continue
                
            # 读取行数据
            for col_idx in range(0, len(field_names)):
                col = col_idx + 1  # Excel列从1开始
                cell_value = sheet.cell(row=row_idx, column=col).value
                field_name = field_names[col_idx]
                field_type = types[col_idx]
                
                converted_value = self.convert_type(cell_value, field_type)
                row_data[field_name] = converted_value
                row_values.append((field_name, converted_value))
            
            # 使用第一列(onlyID)作为唯一键
            row_id = sheet.cell(row=row_idx, column=1).value
            if row_id is not None:
                all_rows.append((str(row_id), row_values))
            else:
                # 如果没有ID，使用行号
                all_rows.append((str(row_idx - 5), row_values))  # 减5是因为从第6行开始
                
        # 写入所有行数据
        for row_id, row_values in all_rows:
            lua_code += f"    [\"{row_id}\"] = {{\n"
            for field_name, value in row_values:
                lua_code += f"        [\"{field_name}\"] = {value},\n"
            lua_code += "    },\n"
            
        lua_code += "}\n\n"
        lua_code += "return " + table_name

        # 写入Lua模块文件
        output_file = Path(self.output_dir) / f"{table_name}.lua"
        with open(output_file, "w", encoding="utf-8") as f:
            f.write(lua_code)
            
        print(f"已生成: {output_file}")
        return table_name
        
    def convert_all(self):
        """转换所有Excel文件"""
        print(f"正在扫描目录: {self.input_dir}")
        processed_files = []
        
        # 扫描所有Excel文件
        for root, _, files in os.walk(self.input_dir):
            for file in files:
                if file.endswith((".xlsx", ".xls")) and not file.startswith("~$"):
                    file_path = os.path.join(root, file)
                    table_name = self.process_excel_file(file_path)
                    processed_files.append(table_name)
                    
        # 生成一个索引模块，用于加载所有数据表
        if processed_files:
            self.generate_index_module(processed_files)
            
        print(f"完成! 共处理了 {len(processed_files)} 个表格。")
        
    def generate_index_module(self, table_names):
        """生成索引模块，方便一次性导入所有表格"""
        lua_code = "-- 表格数据索引模块\n"
        lua_code += "local TableData = {\n"
        
        for table_name in table_names:
            lua_code += f"    {table_name} = require(script.{table_name}),\n"
            
        lua_code += "}\n\n"
        lua_code += "return TableData"
        
        # 写入索引模块
        output_file = Path(self.output_dir) / "TableData.lua"
        with open(output_file, "w", encoding="utf-8") as f:
            f.write(lua_code)
            
        print(f"已生成索引模块: {output_file}")

# 直接在这个脚本中调用转换函数，免除命令行参数解析
if __name__ == "__main__":
    # 当前脚本目录
    script_dir = Path(__file__).parent.absolute()
     
    # 设置输入和输出目录
    input_dir = script_dir / "form"
    output_dir = script_dir.parent / "src" / "Script" / "Outputcfg"
    
    print(f"输入目录: {input_dir}")
    print(f"输出目录: {output_dir}")
    
    # 创建转换器并执行转换
    converter = ExcelToRobloxConverter(input_dir, output_dir)
    converter.convert_all()
    
    print("\n转换完成!")
